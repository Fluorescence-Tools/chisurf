"""TTTR Image Browser Plugin Workspace Widget.

Browse TTTR files in a folder and preview intensity images for all
DetectorWizard-defined detector windows.
"""

from __future__ import annotations

import json
import logging
import os
import pathlib
import tempfile
from typing import Dict, List, Optional, Tuple

import numpy as np
import pyqtgraph as pg
from qtpy.QtCore import QEvent, Qt, QTimer
from qtpy.QtGui import QColor, QFont, QImage, QPainter, QPixmap
from qtpy.QtWidgets import (
    QAbstractItemView,
    QApplication,
    QCheckBox,
    QComboBox,
    QFileDialog,
    QHeaderView,
    QHBoxLayout,
    QLabel,
    QMessageBox,
    QProgressDialog,
    QPushButton,
    QSplitter,
    QTableWidget,
    QTableWidgetItem,
    QTextEdit,
    QVBoxLayout,
    QWidget,
)

from chisurf.gui.widgets.wizard.tttr_channeldefinition import DetectorWizardPage
from chisurf.plugins.tttr.trace_browser.__init__ import NoHoverSelectTable, StarRatingWidget, get_tttr_supported_exts
from chisurf.plugins.tttr.tttr_image_browser.gui.client import TTTRImageBrowserClient

try:
    import tttrlib
except Exception:
    tttrlib = None

try:
    from docx import Document
    from docx.shared import Inches
except Exception:
    Document = None
    Inches = None

try:
    from chisurf.gui.misc_helpers import persist_plugin_state
except ImportError:
    persist_plugin_state = lambda n: lambda c: c

_log = logging.getLogger(__name__)

name = "Imaging:Tools:Image Browser"
META_FILENAME = ".image_browser_meta.json"


@persist_plugin_state("tttr_image_browser")
class TTTRImageBrowser(QWidget):
    """Workspace widget for TTTR Image Browser."""

    def __init__(self, parent=None):
        """Create the workspace layout and components.

        Parameters
        ----------
        parent : QWidget, optional
            The parent widget, by default None.
        """
        super().__init__(parent)
        self.setWindowTitle("TTTR Image Browser")

        # State
        self.current_folder: Optional[pathlib.Path] = None
        self.meta: dict[str, dict] = {}
        self.setup_settings: Optional[dict] = None
        self._current_file: Optional[pathlib.Path] = None
        self._mosaic_cache: dict[pathlib.Path, tuple[np.ndarray, list[str], int, int]] = {}
        self._is_loading = False
        self._meta_save_timer: Optional[QTimer] = None
        self._client = TTTRImageBrowserClient()

        try:
            self._meta_save_timer = QTimer(self)
            self._meta_save_timer.setSingleShot(True)
            self._meta_save_timer.setInterval(300)
            self._meta_save_timer.timeout.connect(self._flush_meta_to_disk)
        except Exception:
            self._meta_save_timer = None

        # Root Layout
        self.root_layout = QVBoxLayout(self)

        # Page 0: Detector Setup
        self.page0 = QWidget(self)
        p0_layout = QVBoxLayout(self.page0)
        p0_layout.addWidget(QLabel("Setup definition (DetectorWizard)", self.page0))
        self.detector_page = DetectorWizardPage(
            show_help=False,
            show_setups_file=True,
            show_setup_selection=True,
            show_tttr_reading=True,
            show_tables=True,
            show_add_inputs=True,
        )
        p0_layout.addWidget(self.detector_page)
        self.btn_continue = QPushButton("Use setup and continue →", self.page0)
        self.btn_continue.clicked.connect(self._on_continue)
        p0_layout.addWidget(self.btn_continue)

        # Page 1: Image Browser
        self.page1 = QWidget(self)
        p1_layout = QVBoxLayout(self.page1)

        # Controls Row
        ctrl_row = QHBoxLayout()
        self.folder_label = QLabel("No folder selected", self.page1)

        self.btn_back = QPushButton("← Back to setup", self.page1)
        self.btn_back.clicked.connect(self._on_back_to_setup)
        ctrl_row.addWidget(self.btn_back)

        self.btn_pick_folder = QPushButton("Pick folder", self.page1)
        self.btn_pick_folder.clicked.connect(self._on_pick_folder)

        self.filter_combo = QComboBox(self.page1)
        self.filter_combo.addItems(["All", "≥ 1★", "≥ 2★★", "≥ 3★★★", "Only 0★"])
        self.filter_combo.currentIndexChanged.connect(self._apply_filter)

        self.btn_clear = QPushButton("Clear", self.page1)
        self.btn_clear.clicked.connect(self._on_clear)

        self.btn_clear_caches = QPushButton("Clear caches", self.page1)
        self.btn_clear_caches.clicked.connect(self._on_clear_caches)

        self.btn_export = QPushButton("Export selected…", self.page1)
        self.btn_export.clicked.connect(self._on_export)

        self.btn_export_docx = QPushButton("Export DOCX…", self.page1)
        self.btn_export_docx.clicked.connect(self._on_export_docx)

        self.btn_save_tiff = QPushButton("Save TIFF…", self.page1)
        self.btn_save_tiff.clicked.connect(self._on_save_tiff)

        ctrl_row.addWidget(self.folder_label)
        ctrl_row.addWidget(self.btn_pick_folder)
        ctrl_row.addWidget(QLabel("Filter:"))
        ctrl_row.addWidget(self.filter_combo)
        ctrl_row.addWidget(self.btn_clear)
        ctrl_row.addWidget(self.btn_clear_caches)

        self.chk_subfolders = QCheckBox("Process subfolders", self.page1)
        self.chk_subfolders.setChecked(False)
        self.chk_subfolders.toggled.connect(self._on_subfolders_toggled)
        ctrl_row.addWidget(self.chk_subfolders)

        ctrl_row.addWidget(self.btn_export)
        ctrl_row.addWidget(self.btn_save_tiff)
        ctrl_row.addWidget(self.btn_export_docx)

        p1_layout.addLayout(ctrl_row)

        # Splitter: left list, right details
        splitter = QSplitter(self.page1)
        splitter.setOrientation(Qt.Horizontal)

        # Left: table of files with rating
        self.table = NoHoverSelectTable(self.page1)
        self.table.setColumnCount(3)
        self.table.setHorizontalHeaderLabels(["File", "Rating", "Size (MB)"])
        self.table.horizontalHeader().setStretchLastSection(False)
        try:
            self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
            self.table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeToContents)
            self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeToContents)
        except Exception:
            self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.ExtendedSelection)
        try:
            self.table.setSortingEnabled(True)
            self.table.horizontalHeader().setSortIndicatorShown(True)
        except Exception:
            pass
        self.table.setStyleSheet("QTableView::item:hover { background: transparent; }")
        self.table.itemSelectionChanged.connect(self._on_selection_changed)
        splitter.addWidget(self.table)

        # Right: image canvas and annotation
        right = QWidget(self.page1)
        right_layout = QVBoxLayout(right)
        self.pg_canvas = pg.GraphicsLayoutWidget(right)
        self.viewbox = self.pg_canvas.addViewBox(lockAspect=True)
        self.viewbox.invertY(False)
        self.image_item = pg.ImageItem()
        self.viewbox.addItem(self.image_item)

        # ColorMap Magma setup
        try:
            if hasattr(self.image_item, "setColorMap"):
                cmap = pg.colormap.get("magma")
                if cmap is not None:
                    self.image_item.setColorMap(cmap)
                else:
                    self._set_fallback_lut()
            else:
                self._set_fallback_lut()
        except Exception:
            pass

        right_layout.addWidget(self.pg_canvas)
        self._overlay_items: list[pg.TextItem] = []
        right_layout.addWidget(QLabel("Annotation:"))
        self.annotation = QTextEdit(self.page1)
        try:
            self.annotation.setMaximumHeight(100)
        except Exception:
            pass
        self.annotation.textChanged.connect(self._on_annotation_changed)
        right_layout.addWidget(self.annotation)
        splitter.addWidget(right)

        p1_layout.addWidget(splitter)

        # Add pages to root
        self.root_layout.addWidget(self.page0)
        self.root_layout.addWidget(self.page1)
        self.page1.hide()

        # Drag-and-drop
        self.setAcceptDrops(True)
        self.table.setAcceptDrops(True)
        self.table.installEventFilter(self)

        # Hide buttons if managed by the toolbar
        for btn in (
            self.btn_pick_folder,
            self.btn_clear,
            self.btn_clear_caches,
            self.btn_export,
            self.btn_save_tiff,
            self.btn_export_docx,
        ):
            try:
                btn.setVisible(False)
            except Exception:
                pass

    def _set_fallback_lut(self):
        from chisurf.plugins.tttr.tttr_image_browser.core.image import get_magma_lut
        lut = get_magma_lut()
        if lut is not None:
            self.image_item.setLookupTable(lut)

    def _allowed_exts_for_setup(self) -> set:
        from chisurf.plugins.tttr.tttr_image_browser.api.io import allowed_exts_for_setup
        return allowed_exts_for_setup(self.setup_settings)

    def _on_continue(self):
        self.setup_settings = self.detector_page.get_settings()
        self.page0.hide()
        self.page1.show()

    def _on_back_to_setup(self):
        _log.info("TTTRImageBrowser: Back to setup")
        self.page1.hide()
        self.page0.show()

    def _on_pick_folder(self):
        folder = QFileDialog.getExistingDirectory(self, "Select folder with TTTR files")
        if not folder:
            return
        self._open_folder(pathlib.Path(folder))

    def _rel_key(self, path: pathlib.Path) -> str:
        try:
            if self.current_folder is not None:
                return str(pathlib.Path(path).resolve().relative_to(self.current_folder.resolve()))
        except Exception:
            pass
        return path.name

    def _meta_get(self, path: pathlib.Path) -> dict:
        key = self._rel_key(path)
        rec = self.meta.get(key)
        if rec is not None:
            return rec
        return self.meta.get(path.name, {})

    def _meta_set(self, path: pathlib.Path, rec: dict):
        key = self._rel_key(path)
        self.meta[key] = rec
        try:
            self._client.set_metadata(str(self.current_folder), dict(self.meta))
        except Exception:
            pass

    def _open_folder(self, folder: pathlib.Path):
        self._is_loading = True
        self.current_folder = folder
        self.folder_label.setText(str(folder))
        _log.info(f"TTTRImageBrowser: Opened folder {folder}")
        self.meta = self._client.get_metadata(str(folder))
        try:
            self._scan_and_fill()
        finally:
            self._is_loading = False

    def _human_size(self, n: int) -> str:
        from chisurf.plugins.tttr.tttr_image_browser.api.io import human_size
        return human_size(n)

    def _scan_and_fill(self):
        self.table.setRowCount(0)
        if self.current_folder is None:
            return
        try:
            recursive = self.chk_subfolders.isChecked()
            rows = self._client.list_files(str(self.current_folder), recursive=recursive, setup_settings=self.setup_settings)
        except Exception as e:
            _log.exception(f"Failed listing files via client: {e}")
            rows = []

        for r_data in rows:
            p = pathlib.Path(r_data["path"])
            rating = r_data["rating"]
            size_bytes = r_data["size"]
            size_text = r_data["size_text"]

            r = self.table.rowCount()
            self.table.insertRow(r)

            try:
                rel_txt = str(p.resolve().relative_to(self.current_folder.resolve()))
            except Exception:
                rel_txt = p.name

            item = QTableWidgetItem(rel_txt)
            item.setData(Qt.UserRole, str(p))
            self.table.setItem(r, 0, item)

            rating_item = QTableWidgetItem()
            rating_item.setFlags(rating_item.flags() & ~Qt.ItemIsEditable)
            rating_item.setData(Qt.EditRole, int(rating))
            self.table.setItem(r, 1, rating_item)

            stars = StarRatingWidget(self.table)
            stars.set_rating(rating)

            def _on_rating_changed(val, path=p, w=stars):
                self._update_rating(path, int(val))
                cur_row = -1
                for rr in range(self.table.rowCount()):
                    if self.table.cellWidget(rr, 1) is w:
                        cur_row = rr
                        break
                if cur_row != -1:
                    it = self.table.item(cur_row, 1)
                    if it is not None:
                        it.setData(Qt.EditRole, int(val))
                self._refresh_list()
                try:
                    header = self.table.horizontalHeader()
                    self.table.sortItems(header.sortIndicatorSection(), header.sortIndicatorOrder())
                except Exception:
                    pass

            stars.ratingChanged.connect(_on_rating_changed)
            self.table.setCellWidget(r, 1, stars)

            size_item = QTableWidgetItem(size_text)
            size_item.setFlags(size_item.flags() & ~Qt.ItemIsEditable)
            size_item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            size_item.setData(Qt.EditRole, float(max(size_bytes, 0)) / (1024.0 * 1024.0))
            self.table.setItem(r, 2, size_item)

        self._apply_filter()
        try:
            self.table.sortItems(0, Qt.AscendingOrder)
        except Exception:
            pass

        try:
            self._precompute_all_images()
        except Exception as _e:
            _log.debug(f"Precompute skipped: {_e}")

    def _update_rating(self, path: pathlib.Path, rating: int):
        rec = self._meta_get(path)
        rec["rating"] = int(rating)
        self._meta_set(path, rec)
        self._schedule_meta_save()
        self._refresh_list()

    def _apply_filter(self):
        self._refresh_list()

    def _schedule_meta_save(self):
        if self._meta_save_timer is not None:
            self._meta_save_timer.start()
        else:
            self._flush_meta_to_disk()

    def _flush_meta_to_disk(self):
        try:
            if self.current_folder:
                from chisurf.plugins.tttr.tttr_image_browser.core.metadata import save_meta
                save_meta(self.current_folder, self.meta)
        except Exception:
            pass

    def _refresh_list(self):
        rows = []
        for r in range(self.table.rowCount()):
            item0 = self.table.item(r, 0)
            if item0 is None:
                continue
            p_str = item0.data(Qt.UserRole)
            if not p_str:
                continue
            p = pathlib.Path(p_str)
            rec = self._meta_get(p)
            rating = int(rec.get("rating", 0))
            rows.append((r, p, rating))

        idx = self.filter_combo.currentIndex()
        allowed_exts = self._allowed_exts_for_setup()

        def accept(rt: int, path: pathlib.Path) -> bool:
            rating_ok = (
                idx == 0
                or (idx == 1 and rt >= 1)
                or (idx == 2 and rt >= 2)
                or (idx == 3 and rt >= 3)
                or (idx == 4 and rt == 0)
            )
            ext_ok = (not allowed_exts) or (path.suffix.lower() in allowed_exts)
            return rating_ok and ext_ok

        rows = [t for t in rows if accept(t[2], t[1])]
        for r in range(self.table.rowCount()):
            self.table.setRowHidden(r, True)
        for r, _, _ in rows:
            self.table.setRowHidden(r, False)

        vis = [r for r, _, _ in rows]
        if vis and not self.table.selectionModel().selectedRows():
            self.table.selectRow(vis[0])

    def _on_clear(self):
        self.table.setRowCount(0)
        self._clear_preview_and_annotation()
        _log.info("TTTRImageBrowser: Cleared file list")

    def _on_selection_changed(self):
        try:
            self._commit_current_annotation()
        except Exception:
            pass
        paths = self._selected_paths()
        if not paths:
            self._clear_preview_and_annotation()
            return
        self._plot_file(paths[0])
        p = paths[0]
        rec = self._meta_get(p)
        self.annotation.blockSignals(True)
        try:
            self.annotation.setPlainText(rec.get("annotation", ""))
        finally:
            self.annotation.blockSignals(False)

    def _on_annotation_changed(self):
        if self._is_loading or not self.current_folder:
            return
        paths = self._selected_paths()
        if not paths:
            return
        p = paths[0]
        rec = self._meta_get(p)
        rec["annotation"] = self.annotation.toPlainText()
        self._meta_set(p, rec)

    def _commit_current_annotation(self):
        if not self.current_folder or not self._current_file:
            return
        rec = self._meta_get(self._current_file)
        rec["annotation"] = self.annotation.toPlainText()
        self._meta_set(self._current_file, rec)

    def _plot_file(self, path: pathlib.Path):
        if path not in self._mosaic_cache:
            res = None
            try:
                res = self._client.load_image(
                    str(path),
                    setup_settings=self.setup_settings,
                    max_side=512,
                    cache_folder=str(self.current_folder) if self.current_folder else None,
                )
            except Exception:
                pass

            if res is not None:
                mosaic = np.array(res["mosaic"], dtype=np.uint8)
                labels = res["labels"]
                cols = res["cols"]
                rows = res["rows"]
                self._mosaic_cache[path] = (mosaic, labels, cols, rows)
            else:
                # Fallback to local computation
                from chisurf.plugins.tttr.tttr_image_browser.core.image import render_mosaic_array
                rendered = render_mosaic_array(path, self._allowed_exts_for_setup(), None)
                if rendered is not None:
                    mosaic, labels, cols, rows, _, _ = rendered
                    self._mosaic_cache[path] = (mosaic, labels, cols, rows)
                else:
                    self._clear_preview_and_annotation()
                    return

        mosaic, labels, cols, rows = self._mosaic_cache[path]
        self.image_item.setImage(mosaic, autoLevels=True)

        for it in self._overlay_items:
            self.viewbox.removeItem(it)
        self._overlay_items = []

        tile_h = mosaic.shape[0] // rows
        tile_w = mosaic.shape[1] // cols
        for idx, text in enumerate(labels):
            r, c = divmod(idx, cols)
            txt = pg.TextItem(text=text, color=(255, 255, 255))
            txt.setAnchor((0, 0))
            txt.setPos(c * tile_w + 3, r * tile_h + 3)
            self.viewbox.addItem(txt)
            self._overlay_items.append(txt)

        self.viewbox.autoRange()
        self._current_file = path

    def _clear_preview_and_annotation(self):
        self.image_item.setImage(np.zeros((1, 1), dtype=np.uint8))
        for it in self._overlay_items:
            try:
                self.viewbox.removeItem(it)
            except Exception:
                pass
        self._overlay_items = []
        self.annotation.clear()
        self._current_file = None

    def _selected_paths(self) -> list[pathlib.Path]:
        sel = []
        for idx in self.table.selectionModel().selectedRows():
            item = self.table.item(idx.row(), 0)
            if item is not None:
                sel.append(pathlib.Path(item.data(Qt.UserRole)))
        return sel

    def _on_subfolders_toggled(self):
        if self.current_folder:
            self._scan_and_fill()

    def _precompute_all_images(self):
        if self.current_folder is None or tttrlib is None:
            return
        paths = []
        for r in range(self.table.rowCount()):
            item = self.table.item(r, 0)
            if item is not None:
                p_str = item.data(Qt.UserRole)
                if p_str:
                    paths.append(pathlib.Path(p_str))
        if not paths:
            return

        files_to_process = [p for p in paths if p not in self._mosaic_cache]
        if not files_to_process:
            return

        dlg = QProgressDialog("Precomputing images...", "Cancel", 0, len(files_to_process), self)
        dlg.setWindowTitle("Precomputing images")
        dlg.setAutoClose(True)
        dlg.setMinimumDuration(0)

        try:
            for i, p in enumerate(files_to_process):
                dlg.setValue(i)
                dlg.setLabelText(f"Processing {p.name} ({i + 1}/{len(files_to_process)})")
                QApplication.processEvents()
                if dlg.wasCanceled():
                    break
                try:
                    res = self._client.load_image(
                        str(p),
                        setup_settings=self.setup_settings,
                        max_side=512,
                        cache_folder=str(self.current_folder),
                    )
                    if res:
                        self._mosaic_cache[p] = (np.array(res["mosaic"], dtype=np.uint8), res["labels"], res["cols"], res["rows"])
                except Exception:
                    pass
        finally:
            dlg.setValue(len(files_to_process))
            dlg.close()

    def eventFilter(self, obj, event):
        if obj is self.table and event is not None:
            et = event.type()
            if et in (QEvent.DragEnter, QEvent.DragMove, QEvent.Drop):
                if et == QEvent.DragEnter:
                    self.dragEnterEvent(event)
                    return True
                elif et == QEvent.DragMove:
                    self.dragMoveEvent(event)
                    return True
                elif et == QEvent.Drop:
                    self.dropEvent(event)
                    return True
            if et == QEvent.KeyPress:
                try:
                    if event.key() in (Qt.Key_Delete,):
                        self._on_delete_selected()
                        return True
                except Exception:
                    pass
        return super().eventFilter(obj, event)

    def _first_dropped_directory(self, event) -> Optional[pathlib.Path]:
        md = event.mimeData()
        if md and md.hasUrls():
            for url in md.urls():
                local = url.toLocalFile()
                if local:
                    p = pathlib.Path(local)
                    if p.exists() and p.is_dir():
                        return p
        return None

    def dragEnterEvent(self, event):
        if self._first_dropped_directory(event) is not None:
            event.acceptProposedAction()
        else:
            event.ignore()

    def dragMoveEvent(self, event):
        if self._first_dropped_directory(event) is not None:
            event.acceptProposedAction()
        else:
            event.ignore()

    def dropEvent(self, event):
        folder = self._first_dropped_directory(event)
        if folder is not None:
            event.acceptProposedAction()
            self._open_folder(folder)
        else:
            event.ignore()

    def _on_delete_selected(self):
        selected_paths = self._selected_paths()
        if not selected_paths or not self.current_folder:
            return
        trash = self.current_folder / ".trash"
        trash.mkdir(parents=True, exist_ok=True)

        import shutil
        import time

        for p in selected_paths:
            try:
                if not p.exists():
                    continue
                dest = trash / p.name
                if dest.exists():
                    ts = time.strftime("%Y%m%d-%H%M%S")
                    dest = dest.with_name(f"{dest.stem}__{ts}{dest.suffix}")
                shutil.move(str(p), str(dest))

                key = self._rel_key(p)
                if key in self.meta:
                    del self.meta[key]
            except Exception as e:
                _log.warning(f"Failed to trash {p}: {e}")

        self._schedule_meta_save()
        self._scan_and_fill()

    def _on_clear_caches(self):
        self._mosaic_cache.clear()
        base = self.current_folder
        if base and base.exists():
            import shutil
            for d in base.rglob(CACHE_DIR_NAME):
                shutil.rmtree(str(d), ignore_errors=True)
            QMessageBox.information(self, "Caches cleared", "Image caches have been cleared.")

    def _on_export(self):
        paths = self._selected_paths()
        if not paths:
            return
        out_dir = QFileDialog.getExistingDirectory(self, "Select destination folder")
        if not out_dir:
            return
        out = pathlib.Path(out_dir)
        out.mkdir(parents=True, exist_ok=True)
        for p in paths:
            try:
                from shutil import copy2
                copy2(str(p), str(out / p.name))
            except Exception as e:
                _log.warning(f"Failed to copy {p}: {e}")

    def _on_save_tiff(self):
        paths = self._selected_paths()
        if not paths:
            return
        out_dir = QFileDialog.getExistingDirectory(self, "Select destination folder for TIFF stacks")
        if not out_dir:
            return
        try:
            res = self._client.export_tiff([str(p) for p in paths], out_dir, self.setup_settings)
            if res:
                _log.info(f"Exported TIFF files: {res}")
        except Exception as e:
            _log.warning(f"TIFF export via client failed: {e}")

    def _on_export_docx(self):
        paths = []
        for r in range(self.table.rowCount()):
            if not self.table.isRowHidden(r):
                item = self.table.item(r, 0)
                if item is not None:
                    paths.append(pathlib.Path(item.data(Qt.UserRole)))
        if not paths or Document is None or not self.current_folder:
            return

        docx_path = self.current_folder / f"{self.current_folder.name or 'images'}.docx"
        doc = Document()
        doc.add_heading("TTTR Image Browser Export", level=1)

        temp_dir = pathlib.Path(tempfile.mkdtemp(prefix="img_docx_"))

        try:
            for p in paths:
                doc.add_heading(p.name, level=2)
                rec = self._meta_get(p)
                doc.add_paragraph(f"Rating: {rec.get('rating', 0)}")
                doc.add_paragraph(f"Annotation: {rec.get('annotation', '')}")

                # Retrieve or render image
                self._plot_file(p)
                if p in self._mosaic_cache:
                    mosaic, _, _, _ = self._mosaic_cache[p]
                    # Convert to QImage and save
                    h, w = mosaic.shape
                    from chisurf.plugins.tttr.tttr_image_browser.core.image import get_magma_lut
                    lut = get_magma_lut()
                    if lut is not None:
                        rgb = lut[mosaic]
                        qimg = QImage(rgb.data, w, h, 3 * w, QImage.Format_RGB888)
                    else:
                        qimg = QImage(mosaic.data, w, h, w, QImage.Format_Grayscale8)

                    img_path = temp_dir / f"{p.stem}.png"
                    if qimg.save(str(img_path), "PNG"):
                        doc.add_picture(str(img_path), width=Inches(6) if Inches else None)

            doc.save(str(docx_path))
            QMessageBox.information(self, "DOCX Export", f"Saved DOCX to {docx_path}")
        except Exception as e:
            _log.exception(f"DOCX export failed: {e}")
        finally:
            import shutil
            shutil.rmtree(str(temp_dir), ignore_errors=True)


if __name__ == "__main__":
    import sys
    app = QApplication(sys.argv)
    w = TTTRImageBrowser()
    w.show()
    sys.exit(app.exec())

# When the plugin is loaded as a module with __name__ == "plugin",
# this code will be executed by the Plugin Manager
if __name__ == "plugin":
    from chisurf.plugins.tttr.tttr_image_browser.gui.tool import TTTRImageBrowserTool
    window = TTTRImageBrowserTool()
    window.show()
    window.raise_()
    window.activateWindow()

