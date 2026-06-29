#!/usr/bin/env python3
"""Download Thorlabs optical component spectra (filters, detectors) into MFDB.

Sources data from publicly available transmission measurements digitised by the
``thor2`` R package (https://github.com/tjconstant/thor2). The original data
was obtained from Thorlabs' published product graphs and datasheets.

Component types included:
  - Longpass edge filters (FEL series)
  - Shortpass edge filters (FES series)
  - Bandpass filters (FB series)
  - Laser-line bandpass filters (FL series)
  - Notch filters (NF series)
  - Neutral-density filters (ND series)
  - Photodetector responsivity (APD120A2)

Usage:
    python -m chisurf.plugins._dev.spectra_downloader.download.thorlabs --db <path>
"""

import io
import logging
import re
import tempfile
import urllib.request
import zipfile
from pathlib import Path

import numpy as np
import openpyxl

from chisurf.plugins._dev.fluorophore_db.mfdb_adapter import (
    DEFAULT_DATABASE_PATH,
    FluorophoreDatabase,
)

logger = logging.getLogger(__name__)

THOR2_DATA_URL = (
    "https://github.com/tjconstant/thor2/archive/refs/heads/master.zip"
)

# Map a data-file prefix → a canonical component kind (resolved via
# COMPONENT_KINDS in mfdb_adapter to its category + default spectrum type).
COMPONENT_KIND_BY_PREFIX: dict[str, str] = {
    "FEL": "longpass",
    "FES": "shortpass",
    "FB": "bandpass",
    "FL": "laserline",
    "NF": "notch",
    "ND": "nd",
}

# Photodetectors (CSV files) → component kind.
DETECTOR_FILES: dict[str, str] = {
    "APD120A2_data.csv": "apd",
}


def _parse_filter_name(filename: str) -> str | None:
    """Extract the canonical Thorlabs part number from a data-file name."""
    stem = Path(filename).stem
    # Remove trailing "_Raw_Data" suffix used for ND files
    stem = stem.replace("_Raw_Data", "")
    return stem.upper()


def _infer_component_kind(name: str) -> str | None:
    # Longer prefixes first so "FES"/"FEL" win over a hypothetical "F".
    for prefix in sorted(COMPONENT_KIND_BY_PREFIX, key=len, reverse=True):
        if name.startswith(prefix):
            return COMPONENT_KIND_BY_PREFIX[prefix]
    return None


def _parse_excel_spectrum(data: bytes):
    """Extract (wavelength_nm, value) arrays from an Excel workbook.

    Tries known sheet names in order, then falls back to the first available
    sheet. Expects wavelength in col C (index 2) and transmission / value in
    col D (index 3), with metadata in rows 1-3.
    """
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    preferred = ["Transmission Data", "Transmission", "Notch Filter"]
    sheet_name = next((s for s in preferred if s in wb.sheetnames), wb.sheetnames[0])
    ws = wb[sheet_name]
    wl_list: list[float] = []
    tr_list: list[float] = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        wl = row[2] if len(row) > 2 else None
        tr = row[3] if len(row) > 3 else None
        if wl is not None and tr is not None:
            try:
                wl_f = float(wl)
                tr_f = float(tr)
                if np.isfinite(wl_f) and np.isfinite(tr_f):
                    wl_list.append(wl_f)
                    tr_list.append(tr_f)
            except (TypeError, ValueError):
                continue
    wb.close()
    if not wl_list:
        return None, None
    arr = np.column_stack([np.array(wl_list), np.array(tr_list)])
    arr = arr[arr[:, 0].argsort()]  # sort ascending
    return arr[:, 0], arr[:, 1]


def _parse_csv_spectrum(data: bytes, sep: str = ";") -> tuple[np.ndarray, np.ndarray] | None:
    """Parse a two-column CSV with optional header."""
    text = data.decode("utf-8", errors="replace")
    wl_list: list[float] = []
    val_list: list[float] = []
    for line in text.strip().splitlines():
        line = line.strip().strip('"')
        if not line or not line[0].isdigit():
            continue
        parts = line.split(sep)
        if len(parts) < 2:
            parts = line.split(",")
        if len(parts) >= 2:
            try:
                wl = float(parts[0].strip().strip('"'))
                val = float(parts[1].strip().strip('"'))
                wl_list.append(wl)
                val_list.append(val)
            except ValueError:
                continue
    if not wl_list:
        return None
    arr = np.column_stack([np.array(wl_list), np.array(val_list)])
    arr = arr[arr[:, 0].argsort()]
    return arr[:, 0], arr[:, 1]


def _optical_properties_from_name(name: str) -> dict[str, str]:
    """Derive optical properties from a Thorlabs part number.

    Returns a dict of (raw) property names → values; ``register_component``
    canonicalizes the keys (e.g. "Center Wavelength (nm)" → ``center_wavelength``).
    """
    # Bandpass / laser-line: FB340-10 → CWL=340 nm, BW=10 nm
    m = re.match(r"FB(\d+)-(\d+)", name)
    if m:
        return {"Center Wavelength (nm)": m.group(1), "Bandwidth (nm)": m.group(2)}
    m = re.match(r"FL(\d+(?:\.\d+)?)-(\d+)", name)
    if m:
        return {"Center Wavelength (nm)": m.group(1), "Bandwidth (nm)": m.group(2)}
    # Longpass: FEL0550 → Cut-On = 550 nm
    m = re.match(r"FEL0*(\d+)", name)
    if m:
        return {"Cut-On Wavelength (nm)": m.group(1)}
    # Shortpass: FES0550 → Cut-Off = 550 nm
    m = re.match(r"FES0*(\d+)", name)
    if m:
        return {"Cut-Off Wavelength (nm)": m.group(1)}
    # Notch: NF533-17 → CWL=533 nm, BW=17 nm
    m = re.match(r"NF(\d+)-(\d+)", name)
    if m:
        return {"Center Wavelength (nm)": m.group(1), "Notch Bandwidth (nm)": m.group(2)}
    # ND: ND01 → OD=0.1
    m = re.match(r"ND(\d+)", name)
    if m:
        return {"Optical Density": f"{int(m.group(1)) / 10.0:.1f}"}
    return {}


def download_thorlabs_to_db(db: FluorophoreDatabase) -> dict[str, int]:
    """Download and import all Thorlabs component spectra.

    Parameters
    ----------
    db : FluorophoreDatabase
        Open database handle.

    Returns
    -------
    dict[str, int]
        Mapping from probe-type name → count of imported items.
    """
    print("Downloading thor2 data archive from GitHub …")
    req = urllib.request.Request(
        THOR2_DATA_URL,
        headers={"User-Agent": "Mozilla/5.0 (ChiSurf)"},
    )
    with urllib.request.urlopen(req) as resp:
        archive_bytes = resp.read()

    counts: dict[str, int] = {}

    with tempfile.TemporaryDirectory() as tmpdir:
        zippath = Path(tmpdir) / "thor2.zip"
        zippath.write_bytes(archive_bytes)

        with zipfile.ZipFile(zippath) as zf:
            # Collect all Excel data files
            xlsx_files = [
                n for n in zf.namelist()
                if n.startswith("thor2-master/data-raw/")
                   and n.endswith(".xlsx")
                   and not Path(n).name.startswith("~$")
            ]
            csv_files = [
                n for n in zf.namelist()
                if n.startswith("thor2-master/data-raw/")
                   and n.endswith(".csv")
            ]

            print(f"Found {len(xlsx_files)} Excel files and {len(csv_files)} CSV files.")

            for name in sorted(xlsx_files):
                stem = _parse_filter_name(Path(name).name)
                if stem is None:
                    continue
                kind = _infer_component_kind(stem)
                if kind is None:
                    continue

                with zf.open(name) as fh:
                    raw = fh.read()

                wl, val = _parse_excel_spectrum(raw)
                if wl is None or len(wl) < 3:
                    print(f"  SKIP {stem}: no valid spectral data")
                    continue

                with db:
                    db.register_component(
                        name=stem,
                        source="thorlabs",
                        kind=kind,
                        source_ref=stem,
                        description=f"Thorlabs {stem}",
                        properties=_optical_properties_from_name(stem),
                        spectra=(wl, val),
                    )

                counts[kind] = counts.get(kind, 0) + 1

            # Import detector CSV files
            for csv_name in sorted(csv_files):
                base = Path(csv_name).name
                if base not in DETECTOR_FILES:
                    continue
                kind = DETECTOR_FILES[base]

                with zf.open(csv_name) as fh:
                    raw = fh.read()

                parsed = _parse_csv_spectrum(raw)
                if parsed is None:
                    print(f"  SKIP {base}: no valid data")
                    continue
                wl, val = parsed

                stem = Path(base).stem.upper().replace("_DATA", "")
                with db:
                    db.register_component(
                        name=stem,
                        source="thorlabs",
                        kind=kind,
                        source_ref=stem,
                        description=f"Thorlabs {stem} detector",
                        spectra=(wl, val),
                    )

                counts[kind] = counts.get(kind, 0) + 1

    print("\nImport summary (by component kind):")
    for kind, cnt in sorted(counts.items()):
        print(f"  {kind:<20s} {cnt:3d} items")
    print(f"\nTotal: {sum(counts.values())} items imported.")
    return counts


def main() -> None:
    """CLI entry point."""
    import argparse

    parser = argparse.ArgumentParser(
        description="Download Thorlabs optical component spectra into MFDB",
    )
    parser.add_argument(
        "--db",
        default=str(DEFAULT_DATABASE_PATH),
        help="MFDB SQLite database path (default: %(default)s)",
    )
    args = parser.parse_args()

    db = FluorophoreDatabase(args.db)
    with db:
        download_thorlabs_to_db(db)
    print("Thorlabs data download complete.")


if __name__ == "__main__":
    main()
